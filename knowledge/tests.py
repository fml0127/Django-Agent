from pathlib import Path
import gzip
import json
import sqlite3
import tempfile
import zipfile
from io import BytesIO, StringIO
from types import SimpleNamespace
from unittest.mock import patch

from django.conf import settings
from django.core.management import call_command
from django.db import connection
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import NoReverseMatch, reverse

from accounts.models import User
from drive import services as drive_services
from content_runtime.converters import ConvertedDocument, DocumentConversionError
from content_runtime.inspectors import inspect_bytes

from .models import ContentExtraction, KBChunk, KBDocument, KnowledgeBase, WikiBuildJob, WikiLink, WikiPage
from . import services, wiki_services
from .sqlite_search import load_sqlite_vec


@override_settings(EMBEDDING_API_KEY="", LLM_API_KEY="", RERANK_API_KEY="")
class KnowledgeTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="alice", password="StrongPass123!")
        self.client.force_login(self.user)

    def _xlsx_bytes(self, sheets):
        from openpyxl import Workbook

        workbook = Workbook()
        default = workbook.active
        workbook.remove(default)
        for sheet_name, rows in sheets.items():
            sheet = workbook.create_sheet(sheet_name)
            for row in rows:
                sheet.append(row)
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        return buffer.getvalue()

    def _pdf_bytes(self, text="Converted office content for RAG."):
        import fitz

        pdf = fitz.open()
        page = pdf.new_page()
        if text:
            page.insert_text((72, 72), text)
        data = pdf.tobytes()
        pdf.close()
        return data

    def test_ingest_text_and_query(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        response = self.client.post(
            reverse("knowledge:ingest", args=[kb.id]),
            {"text": "Django 单体应用包含文件管理、知识库和智能问答。"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(KBChunk.objects.filter(kb=kb).exists())
        chunk = KBChunk.objects.get(kb=kb)
        with connection.cursor() as cursor:
            cursor.execute("SELECT count(*) FROM knowledge_kbchunk_vec WHERE chunk_id = %s", [chunk.id])
            self.assertEqual(cursor.fetchone()[0], 1)
            cursor.execute("SELECT count(*) FROM knowledge_kbchunk_fts WHERE rowid = %s", [chunk.id])
            self.assertEqual(cursor.fetchone()[0], 1)
        response = self.client.get(reverse("knowledge:index"), {"kb": kb.id})
        self.assertContains(response, "去 AI 助手提问")
        self.assertContains(response, f"/assistant/?kb={kb.id}")

    def test_knowledge_page_has_quick_kb_switcher_for_ingest(self):
        kb1 = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        kb2 = KnowledgeBase.objects.create(user=self.user, name="会议资料")

        response = self.client.get(reverse("knowledge:index"), {"kb": kb1.id})

        self.assertContains(response, "当前添加到")
        self.assertContains(response, 'class="kb-switcher"', html=False)
        self.assertContains(response, 'data-autosubmit', html=False)
        self.assertContains(response, f'<option value="{kb1.id}" selected>产品资料', html=False)
        self.assertContains(response, f'<option value="{kb2.id}"', html=False)
        self.assertContains(response, reverse("knowledge:ingest", args=[kb1.id]))

    def test_auto_submit_forms_script_exists(self):
        js = Path(settings.BASE_DIR / "static/js/app.js").read_text(encoding="utf-8")

        self.assertIn("function initAutoSubmitForms", js)
        self.assertIn("form[data-autosubmit]", js)
        self.assertIn("form.requestSubmit()", js)

    def test_sqlite_vec_reload_handles_recreated_raw_connection(self):
        class FakeDjangoConnection:
            vendor = "sqlite"

            def __init__(self):
                self.connection = sqlite3.connect(":memory:")

            def ensure_connection(self):
                if self.connection is None:
                    self.connection = sqlite3.connect(":memory:")

        fake = FakeDjangoConnection()
        try:
            load_sqlite_vec(fake)
            fake._sqlite_vec_loaded = True
            fake.connection.close()
            fake.connection = sqlite3.connect(":memory:")

            load_sqlite_vec(fake)

            self.assertEqual(fake._sqlite_vec_loaded_connection_id, id(fake.connection))
            fake.connection.execute("CREATE VIRTUAL TABLE vec_check USING vec0(id integer primary key, embedding float[4])")
        finally:
            fake.connection.close()

    def test_txt_file_ingest_marks_ready_and_indexes(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile("intro.txt", b"Django RAG uses LangChain loaders.", content_type="text/plain"),
        )

        response = self.client.post(reverse("knowledge:ingest", args=[kb.id]), {"file_ids": [str(user_file.id)]})

        self.assertEqual(response.status_code, 302)
        doc = KBDocument.objects.get(kb=kb, user_file=user_file)
        self.assertEqual(doc.status, KBDocument.STATUS_READY)
        self.assertEqual(doc.detected_family, "text")
        self.assertEqual(doc.extraction.status, ContentExtraction.STATUS_READY)
        self.assertGreater(doc.chunk_count, 0)
        self.assertTrue(KBChunk.objects.filter(document=doc).exists())
        response = self.client.get(reverse("knowledge:index"), {"kb": kb.id})
        self.assertContains(response, "已入库")

    def test_content_inspector_classifies_common_files(self):
        self.assertEqual(inspect_bytes("paper.pdf", sample=b"%PDF-1.7").family, "pdf")
        self.assertEqual(inspect_bytes("image.png", sample=b"\x89PNG\r\n\x1a\n").family, "image")
        self.assertEqual(inspect_bytes("old.doc", sample=b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1").family, "legacy_office")
        self.assertEqual(inspect_bytes("bundle.zip", sample=b"PK\x03\x04").family, "archive")
        self.assertEqual(inspect_bytes("code.java", sample=b"class Demo {}").family, "text")
        self.assertEqual(inspect_bytes("program.f", sample=b"      PROGRAM TEST").family, "text")
        self.assertEqual(inspect_bytes("memo.rtf", sample=b"{\\rtf1\\ansi hello}").family, "rtf")
        self.assertEqual(inspect_bytes("data.xlsx", sample=b"PK\x03\x04").family, "xlsx")
        self.assertEqual(inspect_bytes("notes.txt", sample=b"hello").family, "text")

    def test_non_utf8_text_uses_encoding_fallback(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile("latin.txt", "café document".encode("cp1252"), content_type="text/plain"),
        )

        doc = services.ingest_user_file(kb, user_file)

        self.assertEqual(doc.status, KBDocument.STATUS_READY)
        self.assertTrue(KBChunk.objects.filter(document=doc, content__icontains="café").exists())
        self.assertIn(doc.parser_name, {"TextLoader", "TextEncodingFallback"})
        self.assertEqual(doc.extraction.status, ContentExtraction.STATUS_READY)

    def test_unsupported_file_is_marked_without_vector_rows(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile("archive.bin", b"\x00\x01\x02", content_type="application/octet-stream"),
        )

        response = self.client.post(reverse("knowledge:ingest", args=[kb.id]), {"file_ids": [str(user_file.id)]})

        self.assertEqual(response.status_code, 302)
        doc = KBDocument.objects.get(kb=kb, user_file=user_file)
        self.assertEqual(doc.status, KBDocument.STATUS_UNSUPPORTED)
        self.assertEqual(doc.failure_code, "binary_unsupported")
        self.assertEqual(doc.extraction.status, ContentExtraction.STATUS_UNSUPPORTED)
        self.assertEqual(doc.chunk_count, 0)
        self.assertFalse(KBChunk.objects.filter(document=doc).exists())
        response = self.client.get(reverse("knowledge:index"), {"kb": kb.id})
        self.assertContains(response, "不支持")

    def test_code_file_ingests_as_text(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile(
                "Example.java",
                b"public class Example { String topic = \"RAG pipeline\"; }",
                content_type="text/x-java-source",
            ),
        )

        doc = services.ingest_user_file(kb, user_file)

        self.assertEqual(doc.status, KBDocument.STATUS_READY)
        self.assertEqual(doc.detected_family, "text")
        self.assertTrue(KBChunk.objects.filter(document=doc, content__icontains="RAG pipeline").exists())

    def test_rtf_file_uses_striprtf_extractor(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile(
                "memo.rtf",
                b"{\\rtf1\\ansi This RTF document mentions hybrid retrieval.}",
                content_type="application/rtf",
            ),
        )

        doc = services.ingest_user_file(kb, user_file)

        self.assertEqual(doc.status, KBDocument.STATUS_READY)
        self.assertEqual(doc.detected_family, "rtf")
        self.assertIn(doc.parser_name, {"StripRtfExtractor", "SimpleRtfFallback"})
        self.assertTrue(KBChunk.objects.filter(document=doc, content__icontains="hybrid retrieval").exists())

    def test_xlsx_file_ingests_each_sheet_with_metadata(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        payload = self._xlsx_bytes(
            {
                "Summary": [["Metric", "Value"], ["Hit@6", 0.92], ["MRR", 0.81]],
                "Notes": [["Topic", "Text"], ["RAG", "Hybrid retrieval with FTS5"]],
            }
        )
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile(
                "metrics.xlsx",
                payload,
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        )

        doc = services.ingest_user_file(kb, user_file)

        self.assertEqual(doc.status, KBDocument.STATUS_READY)
        self.assertEqual(doc.detected_family, "xlsx")
        self.assertEqual(doc.parser_name, "SpreadsheetExtractor")
        self.assertEqual(doc.extraction.status, ContentExtraction.STATUS_READY)
        self.assertEqual(doc.parser_metadata["sheet_count"], 2)
        self.assertTrue(KBChunk.objects.filter(document=doc, content__icontains="Hit@6").exists())
        chunk = KBChunk.objects.filter(document=doc, metadata__sheet_name="Summary").first()
        self.assertIsNotNone(chunk)
        self.assertEqual(chunk.metadata["row_count"], 3)
        self.assertEqual(chunk.metadata["column_count"], 2)
        self.assertGreaterEqual(chunk.metadata["non_empty_cells"], 6)
        self.assertIn("cell_range", chunk.metadata)

    def test_empty_xlsx_is_marked_empty_without_index_rows(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        payload = self._xlsx_bytes({"Empty": []})
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile(
                "empty.xlsx",
                payload,
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        )

        doc = services.ingest_user_file(kb, user_file)

        self.assertEqual(doc.status, KBDocument.STATUS_FAILED)
        self.assertEqual(doc.detected_family, "xlsx")
        self.assertEqual(doc.failure_code, "empty_text")
        self.assertFalse(KBChunk.objects.filter(document=doc).exists())

    def test_xlsx_openpyxl_failure_falls_back_to_unstructured_loader(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile(
                "fallback.xlsx",
                b"PK\x03\x04fake",
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        )
        fallback_entry = services.Entry(
            raw="Fallback spreadsheet text",
            content="Fallback spreadsheet text",
            compiled="# fallback\nFallback spreadsheet text",
            title="fallback.xlsx",
            source="fallback.xlsx",
            metadata={"title": "fallback.xlsx", "source": "fallback.xlsx"},
        )
        fallback_result = services.ExtractedEntriesResult(
            entries=[fallback_entry],
            parser_name="UnstructuredExcelLoader",
            method=ContentExtraction.METHOD_LANGCHAIN,
        )

        with patch(
            "knowledge.services._xlsx_entries_with_openpyxl",
            side_effect=services.DocumentParseError("openpyxl failed", parser_name="SpreadsheetExtractor"),
        ), patch("knowledge.services._xlsx_entries_with_langchain", return_value=fallback_result):
            doc = services.ingest_user_file(kb, user_file)

        self.assertEqual(doc.status, KBDocument.STATUS_READY)
        self.assertEqual(doc.parser_name, "UnstructuredExcelLoader")
        self.assertTrue(KBChunk.objects.filter(document=doc, content__icontains="Fallback spreadsheet").exists())

    def test_xlsx_all_extractors_fail_is_isolated(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile(
                "broken.xlsx",
                b"PK\x03\x04fake",
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        )

        with patch(
            "knowledge.services._xlsx_entries_with_openpyxl",
            side_effect=services.DocumentParseError("openpyxl failed", parser_name="SpreadsheetExtractor"),
        ), patch(
            "knowledge.services._xlsx_entries_with_langchain",
            side_effect=RuntimeError("unstructured failed"),
        ):
            doc = services.ingest_user_file(kb, user_file)

        self.assertEqual(doc.status, KBDocument.STATUS_FAILED)
        self.assertEqual(doc.failure_code, "spreadsheet_parse_failed")
        self.assertFalse(KBChunk.objects.filter(document=doc).exists())

    def test_legacy_office_is_marked_as_needs_conversion_without_index_rows(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile("legacy.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1legacy", content_type="application/msword"),
        )

        with patch(
            "knowledge.services.convert_legacy_office_to_pdf",
            side_effect=DocumentConversionError("missing libreoffice", failure_code="needs_conversion"),
        ):
            doc = services.ingest_user_file(kb, user_file)

        self.assertEqual(doc.status, KBDocument.STATUS_UNSUPPORTED)
        self.assertEqual(doc.detected_family, "legacy_office")
        self.assertEqual(doc.failure_code, "needs_conversion")
        self.assertFalse(KBChunk.objects.filter(document=doc).exists())
        self.assertEqual(doc.extraction.failure_code, "needs_conversion")

    def test_legacy_office_conversion_success_enters_pdf_parse_chain(self):
        import fitz

        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "Converted legacy office content for RAG.")
        pdf_bytes = pdf.tobytes()
        pdf.close()

        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile("legacy.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1legacy", content_type="application/msword"),
        )

        with patch(
            "knowledge.services.convert_legacy_office_to_pdf",
            return_value=ConvertedDocument(
                data=pdf_bytes,
                source_family="legacy_office",
                target_format="pdf",
                tool="soffice",
                seconds=0.01,
            ),
        ):
            doc = services.ingest_user_file(kb, user_file)

        self.assertEqual(doc.status, KBDocument.STATUS_READY)
        self.assertEqual(doc.detected_family, "legacy_office")
        self.assertIn("LibreOfficeConverter", doc.parser_name)
        self.assertEqual(doc.parser_metadata["conversion"]["target_format"], "pdf")
        self.assertTrue(KBChunk.objects.filter(document=doc, content__icontains="Converted legacy").exists())

    def test_legacy_office_conversion_failure_is_isolated(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile("legacy.ppt", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1legacy", content_type="application/vnd.ms-powerpoint"),
        )

        with patch(
            "knowledge.services.convert_legacy_office_to_pdf",
            side_effect=DocumentConversionError("bad conversion", failure_code="conversion_failed"),
        ):
            doc = services.ingest_user_file(kb, user_file)

        self.assertEqual(doc.status, KBDocument.STATUS_FAILED)
        self.assertEqual(doc.failure_code, "conversion_failed")
        self.assertFalse(KBChunk.objects.filter(document=doc).exists())

    def test_docx_rich_text_keeps_langchain_loader_path(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile(
                "rich.docx",
                b"PK\x03\x04docx",
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        )
        rich_text = "DOCX rich text " * 30
        original = services.ExtractedEntriesResult(
            entries=[
                services.Entry(
                    raw=rich_text,
                    content=rich_text,
                    compiled=f"# rich.docx\n{rich_text}",
                    title="rich.docx",
                    source="rich.docx",
                    metadata={"title": "rich.docx", "source": "rich.docx"},
                )
            ],
            parser_name="Docx2txtLoader",
            method=ContentExtraction.METHOD_LANGCHAIN,
        )

        with patch("knowledge.services._loader_entries_from_bytes", return_value=original), patch(
            "knowledge.services.convert_office_to_pdf"
        ) as converter:
            doc = services.ingest_user_file(kb, user_file)

        converter.assert_not_called()
        self.assertEqual(doc.status, KBDocument.STATUS_READY)
        self.assertEqual(doc.detected_family, "docx")
        self.assertEqual(doc.parser_name, "Docx2txtLoader")
        self.assertFalse(doc.parser_metadata["text_quality"]["sparse"])

    def test_pptx_sparse_text_triggers_pdf_conversion_fallback(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile(
                "slides.pptx",
                b"PK\x03\x04pptx",
                content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            ),
        )
        sparse = services.ExtractedEntriesResult(
            entries=[
                services.Entry(
                    raw="Hi",
                    content="Hi",
                    compiled="# slides.pptx\nHi",
                    title="slides.pptx",
                    source="slides.pptx",
                    metadata={"title": "slides.pptx", "source": "slides.pptx"},
                )
            ],
            parser_name="UnstructuredPowerPointLoader",
            method=ContentExtraction.METHOD_LANGCHAIN,
        )
        converted = services.ExtractedEntriesResult(
            entries=[
                services.Entry(
                    raw="Converted PPTX PDF content with detailed slide text.",
                    content="Converted PPTX PDF content with detailed slide text.",
                    compiled="# slides.pptx\nConverted PPTX PDF content with detailed slide text.",
                    title="slides.pptx",
                    source="slides.pptx",
                    metadata={"title": "slides.pptx", "source": "slides.pptx"},
                )
            ],
            parser_name="PyMuPDFLoader",
            method=ContentExtraction.METHOD_LANGCHAIN,
        )

        def fake_loader(file_bytes, title, source, profile):
            if profile.family == "pptx":
                return sparse
            if profile.family == "pdf":
                return converted
            raise AssertionError(profile.family)

        with patch("knowledge.services._loader_entries_from_bytes", side_effect=fake_loader), patch(
            "knowledge.services.convert_office_to_pdf",
            return_value=ConvertedDocument(
                data=b"%PDF-1.7\nconverted",
                source_family="pptx",
                target_format="pdf",
                tool="soffice",
                seconds=0.02,
            ),
        ) as converter:
            doc = services.ingest_user_file(kb, user_file)

        converter.assert_called_once()
        self.assertEqual(doc.status, KBDocument.STATUS_READY)
        self.assertEqual(doc.detected_family, "pptx")
        self.assertIn("LibreOfficeConverter+PyMuPDFLoader", doc.parser_name)
        self.assertTrue(doc.parser_metadata["original_text_quality"]["sparse"])
        self.assertTrue(KBChunk.objects.filter(document=doc, content__icontains="Converted PPTX PDF").exists())

    def test_docx_sparse_without_libreoffice_keeps_original_non_empty_text(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile(
                "sparse.docx",
                b"PK\x03\x04docx",
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        )
        sparse = services.ExtractedEntriesResult(
            entries=[
                services.Entry(
                    raw="Short",
                    content="Short",
                    compiled="# sparse.docx\nShort",
                    title="sparse.docx",
                    source="sparse.docx",
                    metadata={"title": "sparse.docx", "source": "sparse.docx"},
                )
            ],
            parser_name="Docx2txtLoader",
            method=ContentExtraction.METHOD_LANGCHAIN,
        )

        with patch("knowledge.services._loader_entries_from_bytes", return_value=sparse), patch(
            "knowledge.services.convert_office_to_pdf",
            side_effect=DocumentConversionError("missing libreoffice", failure_code="needs_conversion"),
        ):
            doc = services.ingest_user_file(kb, user_file)

        self.assertEqual(doc.status, KBDocument.STATUS_READY)
        self.assertEqual(doc.parser_name, "Docx2txtLoader")
        self.assertTrue(doc.parser_metadata["text_quality"]["sparse"])
        self.assertEqual(doc.parser_metadata["sparse_fallback"]["failure_code"], "needs_conversion")
        self.assertTrue(KBChunk.objects.filter(document=doc, content__icontains="Short").exists())

    @override_settings(VISION_API_KEY="")
    def test_docx_empty_after_pdf_conversion_marks_vision_not_configured(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile(
                "empty.docx",
                b"PK\x03\x04docx",
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        )

        real_loader = services._loader_entries_from_bytes

        def fake_loader(file_bytes, title, source, profile):
            if profile.family == "docx":
                raise services.DocumentParseError(
                    "empty docx",
                    failure_code="empty_text",
                    profile=profile,
                    parser_name="Docx2txtLoader",
                )
            return real_loader(file_bytes, title, source, profile)

        with patch("knowledge.services._loader_entries_from_bytes", side_effect=fake_loader), patch(
            "knowledge.services.convert_office_to_pdf",
            return_value=ConvertedDocument(
                data=self._pdf_bytes(""),
                source_family="docx",
                target_format="pdf",
                tool="soffice",
                seconds=0.02,
            ),
        ):
            doc = services.ingest_user_file(kb, user_file)

        self.assertEqual(doc.status, KBDocument.STATUS_FAILED)
        self.assertEqual(doc.failure_code, "vision_not_configured")
        self.assertFalse(KBChunk.objects.filter(document=doc).exists())

    def test_gzip_archive_is_decompressed_and_indexed(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        payload = gzip.compress(b"Archive text explains local knowledge indexing.")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile("notes.txt.gz", payload, content_type="application/gzip"),
        )

        doc = services.ingest_user_file(kb, user_file)

        self.assertEqual(doc.status, KBDocument.STATUS_READY)
        self.assertEqual(doc.detected_family, "archive")
        self.assertIn("ArchiveExtractor", doc.parser_name)
        self.assertTrue(KBChunk.objects.filter(document=doc, content__icontains="local knowledge indexing").exists())

    def test_zip_archive_is_decompressed_and_indexed(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        buffer = BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("inside/readme.txt", "Zip member contains searchable RAG content.")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile("bundle.zip", buffer.getvalue(), content_type="application/zip"),
        )

        doc = services.ingest_user_file(kb, user_file)

        self.assertEqual(doc.status, KBDocument.STATUS_READY)
        self.assertEqual(doc.detected_family, "archive")
        self.assertTrue(KBChunk.objects.filter(document=doc, content__icontains="searchable RAG").exists())

    def test_zip_archive_without_supported_members_is_isolated(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        buffer = BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("payload.bin", b"\x00\x01\x02")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile("bundle.zip", buffer.getvalue(), content_type="application/zip"),
        )

        doc = services.ingest_user_file(kb, user_file)

        self.assertEqual(doc.status, KBDocument.STATUS_FAILED)
        self.assertEqual(doc.failure_code, "archive_no_supported_files")
        self.assertFalse(KBChunk.objects.filter(document=doc).exists())

    def test_zip_archive_path_traversal_is_rejected(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        buffer = BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("../evil.txt", "unsafe")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile("unsafe.zip", buffer.getvalue(), content_type="application/zip"),
        )

        doc = services.ingest_user_file(kb, user_file)

        self.assertEqual(doc.status, KBDocument.STATUS_FAILED)
        self.assertEqual(doc.failure_code, "archive_limit_exceeded")
        self.assertFalse(KBChunk.objects.filter(document=doc).exists())

    def test_image_file_uses_mocked_vision_extraction_and_indexes(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile("scan.png", b"\x89PNG\r\n\x1a\nfake", content_type="image/png"),
        )

        with patch("knowledge.services.extract_image_file", return_value="# 扫描件\n图片里有知识库内容。"):
            doc = services.ingest_user_file(kb, user_file)

        self.assertEqual(doc.status, KBDocument.STATUS_READY)
        self.assertEqual(doc.detected_family, "image")
        self.assertEqual(doc.parser_name, "VisionImageExtractor")
        self.assertEqual(doc.extraction.method, ContentExtraction.METHOD_VISION)
        self.assertTrue(KBChunk.objects.filter(document=doc, content__icontains="知识库内容").exists())

    @override_settings(VISION_API_KEY="")
    def test_image_file_without_vision_key_fails_without_index_rows(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile("scan.jpg", b"\xff\xd8\xfffake", content_type="image/jpeg"),
        )

        doc = services.ingest_user_file(kb, user_file)

        self.assertEqual(doc.status, KBDocument.STATUS_FAILED)
        self.assertEqual(doc.failure_code, "vision_not_configured")
        self.assertEqual(doc.extraction.failure_code, "vision_not_configured")
        self.assertFalse(KBChunk.objects.filter(document=doc).exists())

    def test_reingesting_same_file_replaces_existing_document(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        user_file = drive_services.save_uploaded_file(
            self.user,
            SimpleUploadedFile("intro.md", b"# Title\nRepeated import should replace old chunks.", content_type="text/markdown"),
        )

        services.ingest_user_file(kb, user_file)
        first_doc_id = KBDocument.objects.get(kb=kb, user_file=user_file).id
        services.ingest_user_file(kb, user_file)

        docs = KBDocument.objects.filter(kb=kb, user_file=user_file)
        self.assertEqual(docs.count(), 1)
        self.assertNotEqual(docs.get().id, first_doc_id)
        self.assertEqual(ContentExtraction.objects.filter(kb=kb, user_file=user_file).count(), 1)
        self.assertEqual(kb.documents.filter(status=KBDocument.STATUS_READY).count(), 1)

    def test_fts_trigram_and_delete_cleanup(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        doc = services.ingest_text(kb, "text", "manual", "手动文本", "知识库可以检索中文片段和文件管理能力。")
        chunk = KBChunk.objects.get(document=doc)

        fts_hits = services.fts_candidates(kb, "中文片段", 6)
        self.assertEqual(fts_hits[0][0], chunk.id)

        hits = services.search(kb, "怎么检索中文片段", top_k=6)
        self.assertTrue(any(hit_chunk.id == chunk.id for _, hit_chunk in hits))

        doc.delete()
        with connection.cursor() as cursor:
            cursor.execute("SELECT count(*) FROM knowledge_kbchunk_vec WHERE chunk_id = %s", [chunk.id])
            self.assertEqual(cursor.fetchone()[0], 0)
            cursor.execute("SELECT count(*) FROM knowledge_kbchunk_fts WHERE rowid = %s", [chunk.id])
            self.assertEqual(cursor.fetchone()[0], 0)

    def test_search_with_trace_preserves_search_compatibility(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        services.ingest_text(kb, "text", "manual", "检索文档", "SQLite FTS5 和 sqlite-vec 可以做混合检索。")

        traced = services.search_with_trace(kb, "混合检索", top_k=3)
        plain_hits = services.search(kb, "混合检索", top_k=3)

        self.assertTrue(traced["hits"])
        self.assertEqual([hit.chunk_id if hasattr(hit, "chunk_id") else hit.chunk.id for hit in traced["hits"]], [hit.chunk.id for hit in plain_hits])
        self.assertIn("rewritten_queries", traced["trace"])
        self.assertIn("vector_candidates", traced["trace"])
        self.assertIn("fts_candidates", traced["trace"])
        self.assertIn("fusion_candidates", traced["trace"])
        self.assertIn("rerank", traced["trace"])
        self.assertIn("final_hits", traced["trace"])

    def test_evaluate_rag_command_outputs_metrics(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        services.ingest_text(kb, "text", "manual", "SQLite 检索", "SQLite FTS5 和 sqlite-vec 可以一起做混合检索。")
        case = {
            "question": "什么可以做混合检索？",
            "expected_document_title": "SQLite 检索",
            "expected_contains": "sqlite-vec",
        }
        with tempfile.NamedTemporaryFile("w+", encoding="utf-8", suffix=".jsonl") as dataset:
            dataset.write(json.dumps(case, ensure_ascii=False) + "\n")
            dataset.flush()
            stdout = StringIO()
            call_command("evaluate_rag", "--kb", str(kb.id), "--dataset", dataset.name, stdout=stdout)

        payload = json.loads(stdout.getvalue())
        self.assertEqual(payload["case_count"], 1)
        self.assertEqual(payload["hit_at_k"], 1.0)
        self.assertEqual(payload["mrr"], 1.0)
        self.assertIn("stage_metrics", payload)
        self.assertIn("vector", payload["stage_metrics"])
        self.assertIn("fts", payload["stage_metrics"])
        self.assertIn("fusion", payload["stage_metrics"])
        self.assertIn("final", payload["stage_metrics"])
        self.assertIn("vector_hit_at_k", payload)
        self.assertIn("vector_hit@k", payload)
        self.assertIn("final_mrr", payload)
        self.assertEqual(payload["stage_metrics"]["final"]["hit_at_k"], 1.0)
        self.assertIn("rerank_analysis", payload)
        self.assertIn("stage_results", payload["results"][0])
        self.assertIn("vector", payload["results"][0]["stage_results"])
        self.assertIn("fts", payload["results"][0]["stage_results"])
        self.assertIn("trace", payload["results"][0])
        self.assertTrue(payload["results"][0]["final_hits"])

    def test_evaluate_rag_command_markdown_and_save_json(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        services.ingest_text(kb, "text", "manual", "SQLite 检索", "SQLite FTS5 和 sqlite-vec 可以一起做混合检索。")
        case = {
            "question": "sqlite-vec 有什么作用？",
            "expected_document_title": "SQLite 检索",
            "expected_contains": "sqlite-vec",
        }
        with tempfile.NamedTemporaryFile("w+", encoding="utf-8", suffix=".jsonl") as dataset:
            dataset.write(json.dumps(case, ensure_ascii=False) + "\n")
            dataset.flush()
            with tempfile.TemporaryDirectory() as output_dir:
                saved_json = Path(output_dir) / "rag_eval_result.json"
                stdout = StringIO()
                call_command(
                    "evaluate_rag",
                    "--kb",
                    str(kb.id),
                    "--dataset",
                    dataset.name,
                    "--format",
                    "markdown",
                    "--save-json",
                    str(saved_json),
                    stdout=stdout,
                )

                markdown = stdout.getvalue()
                self.assertIn("## Summary", markdown)
                self.assertIn("## Stage Comparison", markdown)
                self.assertIn("## Rerank Analysis", markdown)
                self.assertIn("## Failed Cases", markdown)
                self.assertIn("## Case Details", markdown)
                saved_payload = json.loads(saved_json.read_text(encoding="utf-8"))
                self.assertEqual(saved_payload["case_count"], 1)
                self.assertIn("stage_metrics", saved_payload)
                self.assertTrue(saved_payload["results"][0]["stage_results"]["final"]["hit"])

    def test_knowledge_page_does_not_expose_mindmap(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        response = self.client.get(reverse("knowledge:index"), {"kb": kb.id})

        self.assertNotContains(response, "思维导图")
        self.assertNotContains(response, "mindmap")
        self.assertNotContains(response, "向知识库提问")
        with self.assertRaises(NoReverseMatch):
            reverse("knowledge:mindmap", args=[kb.id])
        with self.assertRaises(NoReverseMatch):
            reverse("knowledge:query", args=[kb.id])

    def _fake_wiki_openai(self):
        class FakeChatCompletions:
            def create(self, **kwargs):
                prompt = kwargs["messages"][1]["content"]
                if "source 页面摘要" in prompt:
                    content = (
                        "## Topic Summary\n"
                        "当前知识库整理了 Django、RAG 和 Wiki 的关系。\n\n"
                        "## Major Sources\n"
                        "- [[手动文本]]\n\n"
                        "## Key Conclusions\n"
                        "- Wiki 适合沉淀结构化知识。\n\n"
                        "## Gaps\n"
                        "- 还缺少运行指标。"
                    )
                else:
                    content = (
                        "## Summary\n"
                        "手动文本说明 Django RAG 可以生成 Wiki 页面。\n\n"
                        "## Key Points\n"
                        "- Wiki source 页面来自已入库文档。\n"
                        "- 检索会同时使用 Wiki 和 chunk。\n\n"
                        "## Useful Quotes\n"
                        "- Django RAG uses Wiki.\n\n"
                        "## Connections\n"
                        "- [[MissingConcept]]\n\n"
                        "## Open Questions\n"
                        "- 是否需要图谱。"
                    )
                return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=content))])

        class FakeOpenAI:
            def __init__(self, **kwargs):
                self.chat = SimpleNamespace(completions=FakeChatCompletions())

        return FakeOpenAI

    @override_settings(LLM_API_KEY="test-key", LLM_BASE_URL="https://example.com/v1", LLM_MODEL="test-model")
    def test_build_wiki_generates_source_overview_and_indexes(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        doc = services.ingest_text(kb, "text", "manual", "手动文本", "Django RAG uses Wiki pages.")

        with patch("knowledge.wiki_services.OpenAI", self._fake_wiki_openai()):
            job = wiki_services.build_wiki(kb)

        self.assertEqual(job.status, WikiBuildJob.STATUS_SUCCESS)
        source = WikiPage.objects.get(kb=kb, page_type=WikiPage.TYPE_SOURCE, source_document=doc)
        overview = WikiPage.objects.get(kb=kb, page_type=WikiPage.TYPE_OVERVIEW)
        self.assertEqual(source.status, WikiPage.STATUS_READY)
        self.assertEqual(overview.status, WikiPage.STATUS_READY)
        with connection.cursor() as cursor:
            cursor.execute("SELECT count(*) FROM knowledge_wikipage_vec WHERE page_id = %s", [source.id])
            self.assertEqual(cursor.fetchone()[0], 1)
            cursor.execute("SELECT count(*) FROM knowledge_wikipage_fts WHERE rowid = %s", [source.id])
            self.assertEqual(cursor.fetchone()[0], 1)
        self.assertTrue(WikiLink.objects.filter(source_page=source, target_title="MissingConcept").exists())
        health = wiki_services.wiki_health(kb)
        self.assertFalse(health["missing_overview"])
        self.assertTrue(health["broken_links"])
        self.assertIn("orphan_pages", health)
        self.assertIn("sparse_pages", health)
        self.assertIn("link_density", health)

        response = self.client.get(reverse("knowledge:index"), {"kb": kb.id})
        self.assertContains(response, "Wiki")
        self.assertContains(response, "打开总览")
        self.assertContains(response, "MissingConcept")
        self.assertContains(response, "关系图数据")

        graph_response = self.client.get(reverse("knowledge:wiki_graph_json", args=[kb.id]))
        self.assertEqual(graph_response.status_code, 200)
        graph = graph_response.json()
        self.assertEqual(graph["kb"]["id"], kb.id)
        self.assertTrue(graph["nodes"])
        self.assertIn("edges", graph)
        self.assertIn("health", graph)
        self.assertIn("orphan_page_count", graph["health"])

    def test_wiki_graph_json_rejects_other_users_kb(self):
        other = User.objects.create_user(username="bob", password="StrongPass123!")
        kb = KnowledgeBase.objects.create(user=other, name="私有资料")

        response = self.client.get(reverse("knowledge:wiki_graph_json", args=[kb.id]))

        self.assertEqual(response.status_code, 404)

    @override_settings(LLM_API_KEY="test-key", LLM_BASE_URL="https://example.com/v1", LLM_MODEL="test-model")
    def test_build_wiki_refresh_does_not_duplicate_pages(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        services.ingest_text(kb, "text", "manual", "手动文本", "重复生成不应该产生重复 Wiki 页面。")

        with patch("knowledge.wiki_services.OpenAI", self._fake_wiki_openai()):
            wiki_services.build_wiki(kb)
            wiki_services.build_wiki(kb)

        self.assertEqual(WikiPage.objects.filter(kb=kb, page_type=WikiPage.TYPE_SOURCE).count(), 1)
        self.assertEqual(WikiPage.objects.filter(kb=kb, page_type=WikiPage.TYPE_OVERVIEW).count(), 1)

    @override_settings(LLM_API_KEY="")
    def test_build_wiki_without_llm_key_fails_visibly(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        services.ingest_text(kb, "text", "manual", "手动文本", "Wiki 生成需要 LLM。")

        job = wiki_services.build_wiki(kb)

        self.assertEqual(job.status, WikiBuildJob.STATUS_FAILED)
        self.assertIn("LLM_API_KEY", job.error_message)
        self.assertFalse(WikiPage.objects.filter(kb=kb).exists())

    @override_settings(LLM_API_KEY="test-key", LLM_BASE_URL="https://example.com/v1", LLM_MODEL="test-model")
    def test_deleting_document_marks_source_wiki_page_stale(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        doc = services.ingest_text(kb, "text", "manual", "手动文本", "删除文档后 Wiki 应标记过期。")

        with patch("knowledge.wiki_services.OpenAI", self._fake_wiki_openai()):
            wiki_services.build_wiki(kb)
        page = WikiPage.objects.get(kb=kb, page_type=WikiPage.TYPE_SOURCE)

        doc.delete()
        page.refresh_from_db()

        self.assertEqual(page.status, WikiPage.STATUS_STALE)
        self.assertIsNone(page.source_document)

    @override_settings(LLM_API_KEY="test-key", LLM_BASE_URL="https://example.com/v1", LLM_MODEL="test-model")
    def test_wiki_page_view_renders_markdown_and_enforces_owner(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        services.ingest_text(kb, "text", "manual", "手动文本", "Wiki 页面可以渲染 Markdown。")

        with patch("knowledge.wiki_services.OpenAI", self._fake_wiki_openai()):
            wiki_services.build_wiki(kb)
        page = WikiPage.objects.get(kb=kb, page_type=WikiPage.TYPE_SOURCE)

        response = self.client.get(reverse("knowledge:wiki_page", args=[kb.id, page.slug]))
        self.assertContains(response, "<h2>Summary</h2>", html=True)
        self.assertContains(response, "MissingConcept")

        other = User.objects.create_user(username="bob", password="StrongPass123!")
        other_kb = KnowledgeBase.objects.create(user=other, name="私有资料")
        other_page = WikiPage.objects.create(
            kb=other_kb,
            page_type=WikiPage.TYPE_OVERVIEW,
            slug="overview",
            title="私有总览",
            content="## Summary\nsecret",
            status=WikiPage.STATUS_READY,
        )
        response = self.client.get(reverse("knowledge:wiki_page", args=[other_kb.id, other_page.slug]))
        self.assertEqual(response.status_code, 404)

    @override_settings(LLM_API_KEY="test-key", LLM_BASE_URL="https://example.com/v1", LLM_MODEL="test-model")
    def test_query_rewrite_uses_llm_json(self):
        class FakeChatCompletions:
            def create(self, **kwargs):
                return SimpleNamespace(
                    choices=[SimpleNamespace(message=SimpleNamespace(content='{"queries":["中文检索", "文件管理"]}'))]
                )

        class FakeOpenAI:
            def __init__(self, **kwargs):
                self.chat = SimpleNamespace(completions=FakeChatCompletions())

        with patch("knowledge.services.OpenAI", FakeOpenAI):
            self.assertEqual(services.rewrite_rag_queries("怎么查资料"), ["中文检索", "文件管理"])

    @override_settings(RERANK_API_KEY="test-key", RERANK_MODEL="qwen3-vl-rerank")
    def test_rerank_orders_candidates_by_relevance_score(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        doc1 = services.ingest_text(kb, "text", "manual", "文档一", "第一段讲文件上传。")
        doc2 = services.ingest_text(kb, "text", "manual", "文档二", "第二段讲向量检索。")
        chunk1 = KBChunk.objects.get(document=doc1)
        chunk2 = KBChunk.objects.get(document=doc2)

        fake_response = SimpleNamespace(
            raise_for_status=lambda: None,
            json=lambda: {
                "output": {
                    "results": [
                        {"index": 1, "relevance_score": 0.95},
                        {"index": 0, "relevance_score": 0.12},
                    ]
                }
            },
        )

        with patch("knowledge.services.vector_candidates", return_value=[(chunk1.id, 0.1), (chunk2.id, 0.2)]), patch(
            "knowledge.services.fts_candidates", return_value=[]
        ), patch("knowledge.services.httpx.Client") as client_class:
            client = client_class.return_value.__enter__.return_value
            client.post.return_value = fake_response
            hits = services.search(kb, "怎么做向量检索", top_k=2)

        self.assertEqual(hits[0].chunk.id, chunk2.id)
        self.assertEqual(hits[0].rerank_score, 0.95)

    @override_settings(RERANK_API_KEY="test-key", RERANK_MODEL="qwen3-vl-rerank")
    def test_rerank_failure_keeps_fused_order(self):
        kb = KnowledgeBase.objects.create(user=self.user, name="产品资料")
        doc1 = services.ingest_text(kb, "text", "manual", "文档一", "第一段讲文件上传。")
        doc2 = services.ingest_text(kb, "text", "manual", "文档二", "第二段讲向量检索。")
        chunk1 = KBChunk.objects.get(document=doc1)
        chunk2 = KBChunk.objects.get(document=doc2)

        with patch("knowledge.services.vector_candidates", return_value=[(chunk1.id, 0.1), (chunk2.id, 0.2)]), patch(
            "knowledge.services.fts_candidates", return_value=[]
        ), patch("knowledge.services.httpx.Client") as client_class:
            client_class.return_value.__enter__.return_value.post.side_effect = RuntimeError("down")
            hits = services.search(kb, "怎么做向量检索", top_k=2)

        self.assertEqual(hits[0].chunk.id, chunk1.id)

    # Create your tests here.
